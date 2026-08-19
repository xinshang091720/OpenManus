from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\jly23\Desktop\OpenManus-main")
FRAME_DIR = ROOT / "output" / "wechat_video_plan"
OUTPUT = Path(r"C:\Users\jly23\Desktop\CAD-Revit-IFC_微信公众号90秒视频制作规划_高级技术版.docx")


def font(run, size=10.5, bold=False, color="1F2937"):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    tc_pr.append(node)


def table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    w = tbl_pr.find(qn("w:tblW"))
    if w is None:
        w = OxmlElement("w:tblW")
        tbl_pr.append(w)
    w.set(qn("w:w"), str(sum(widths)))
    w.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    tbl_pr.append(indent)
    for grid, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def cell_text(cell, text, *, size=9, bold=False, color="1F2937"):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    font(p.add_run(text), size=size, bold=bold, color=color)


def para(doc, text, *, size=10.5, after=6, color="1F2937", bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(8)
    font(p.add_run(text), size=16, bold=True, color="1F4D78")


def add_bullets(doc, values):
    for value in values:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.18
        font(p.add_run(value), size=9.6)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(header.add_run("AI 智建 | 微信公众号视频制作规划（高级技术版）"), size=8.5, color="697586")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run("模型交付 Agent｜剪辑与配音执行稿"), size=8, color="697586")

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    font(title.add_run("CAD-Revit-IFC 智能交付闭环"), size=25, bold=True, color="123B5D")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(13)
    font(sub.add_run("微信公众号短视频前期制作规划｜高级技术版｜成片建议 90 秒"), size=12, color="536779")

    strip = doc.add_table(rows=1, cols=3)
    table_geometry(strip, [3120, 3120, 3120])
    for cell, label, value in zip(strip.rows[0].cells,
                                  ["智能体架构", "数据与算法", "最终交付"],
                                  ["Agent → MCP Skills → Revit API", "工程知识检索 + 多层语义匹配 + 结构化决策", "RVT · IFC · XLSX · DOCX"]):
        shade(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        font(p.add_run(label + "\n"), size=8.4, bold=True, color="1F4D78")
        font(p.add_run(value), size=9.5)

    h1(doc, "本版表达原则")
    para(doc, "旁白以 Agent 架构为主线：语义推理模块理解工程语料，MCP Skills 选择专业能力，本地 Revit API 负责确定性执行，标准 IFC 知识库输出合规语义。每一屏只呈现一个清晰的算法或系统能力。", color="1F4D78")
    para(doc, "参考表达来源：用户提供的 PKPM 公众号视频出处：https://mp.weixin.qq.com/s/A52FrOqTJDQXKd_Cq4Cdng", size=8.8, color="536779")

    h1(doc, "90 秒高级技术旁白与分镜")
    table = doc.add_table(rows=1, cols=4)
    table_geometry(table, [820, 2180, 3900, 2460])
    for cell, value in zip(table.rows[0].cells, ["时间", "画面", "旁白", "屏幕技术标签"]):
        shade(cell, "DDE8F2")
        cell_text(cell, value, size=9, bold=True, color="1F4D78")
    rows = [
        ("0–07s", "DWG、RVT、IFC 文件与模型画面快速切换。", "这是一个基于 Transformer 架构的 BIM Agent：它通过 MCP Skills 编排 CAD 解析、Revit 建模、IFC 赋参和报建质检，形成端到端的自动化交付闭环。", "Transformer Agent｜MCP Skills"),
        ("07–22s", "CAD 与 Revit 对照；放大图纸文字和轴网。", "Agent 先把天正 DWG 转为结构化 DXF，递归提取 TEXT、MTEXT、嵌套块和属性文字；再由语义推理模块将图纸文本映射为房间功能、空间名称和坐标语义。", "DXF Entity Parsing｜Semantic Reasoning"),
        ("22–34s", "轴网局部、执行前模型，转入房间结果。", "系统以 CAD 与 Revit 同名轴网为几何锚点，构建坐标配准关系，计算轴网中点的平移向量，并将空间语义精准投射到 Revit 模型坐标系。", "Geometric Alignment｜Translation Vector"),
        ("34–45s", "执行后彩色模型与属性栏；局部拉远。", "房间创建 Skill 调用本地 Revit API，批量生成闭合区域房间，按楼层写入名称与坐标结果，并自动沉淀版本化模型和审计数据。", "MCP Skill Execution｜BatchCreateRooms"),
        ("45–60s", "切入 IFC 参数栏、工作流结果卡片。", "进入 IFC 阶段，Agent 连接深圳 BIM 标准标识库，以 standardId 109003 检索专业候选标识；标高与各专业构件进入对应的知识目录。", "IFC Knowledge Base｜Standard 109003"),
        ("60–72s", "参数栏放大，随后显示赋参结果。", "系统把构件类型、族和类别组织成三级语义路径，完成候选召回和加权排序；决策模块在 Top-K 范围内输出结构化 JSON，再批量写回 IFC 参数。", "Semantic Graph｜Top-K Retrieval｜JSON"),
        ("72–82s", "导出阶段与文件路径动画；快速切换。", "已赋参模型进入 ifc-assigned 交付检查点。系统导出 IFC 与配套表格，并以文件指纹建立本次交付物的可追溯索引。", "Artifact Index｜IFC + XLSX"),
        ("82–90s", "SZ-IFC 检查结果和 Word 报告收束。", "最后，Agent 通过 Windows UI Automation 驱动 SZ-IFC，加载专业规则、执行模型检查并生成 Word 质检报告。图纸、模型、标准数据和质检成果由同一 Agent 统一闭环。", "UI Automation｜Rule Engine｜DOCX"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, row)):
            if index == 0:
                shade(cell, "F4F6F9")
                cell_text(cell, value, size=8.8, bold=True, color="1F4D78")
            elif index == 3:
                cell_text(cell, value, size=8.5, bold=True, color="536779")
            else:
                cell_text(cell, value, size=8.8)

    doc.add_page_break()
    h1(doc, "AI 核心能力矩阵")
    para(doc, "该段用于主播、剪辑与答疑统一口径。以“Agent 架构—语义理解—工程知识—自动化执行”构成技术叙事主线，所有关键词都与工作流的真实输入、接口和输出相对应。", size=9.5, color="536779")
    facts = doc.add_table(rows=1, cols=3)
    table_geometry(facts, [1920, 3600, 3840])
    for cell, text in zip(facts.rows[0].cells, ["AI 能力层", "核心算法 / 架构", "视频可视化方式"]):
        shade(cell, "DDE8F2")
        cell_text(cell, text, size=9, bold=True, color="1F4D78")
    facts_rows = [
        ("Agent 编排", "Agent 进行任务理解；MCP Skills 将 CAD、Revit、IFC 与 SZ-IFC 能力按业务阶段编排为可执行链路。", "画面标签：Agent Orchestration。"),
        ("工程语义理解", "天正对象预处理后转 DXF；递归解析图层、TEXT/MTEXT、ATTRIB 与嵌套块。语义模型以受控 JSON 完成空间语义判定。", "画面标签：DXF 图元解析 + 语义推理。"),
        ("坐标智能", "以 CAD-Revit 轴网为几何锚点，执行尺度与方向特征校验，并计算中点平移向量，实现空间文本到 BIM 坐标系的映射。", "轴网连线 + ΔX / ΔY 动效。"),
        ("标准知识库", "本地 Revit 插件通过 GetIfcIdent / GetLevelIfcIdent 检索深圳标准 109003 的 IFC 标识目录，建立专业化候选集合。", "画面标签：IFC Knowledge Base。"),
        ("语义匹配", "构件类型、族、类别形成三级语义图；父级路径回溯、精确/包含加权排序与 Top-K 结构化复核共同完成标识决策。", "候选 1–5 列表，选中最终标识。"),
        ("自动化执行", "Python 异步工作流经本地 HTTP 插件驱动 Revit；SZ-IFC 由 Windows UI Automation 完成规则检查与 DOCX 报告输出。", "画面标签：MCP Skill + UI Automation。"),
        ("语义增强", "BERT / Embedding 编码可作为工程文本与标准词库的语义索引层，持续增强跨项目术语归一与标准检索能力。", "画面标签：BERT Embedding｜Semantic Index。"),
    ]
    for row in facts_rows:
        cells = facts.add_row().cells
        for index, (cell, text) in enumerate(zip(cells, row)):
            cell_text(cell, text, size=8.7, bold=(index == 0), color="1F4D78" if index == 0 else "1F2937")

    h1(doc, "关键画面与镜头动作")
    frame_table = doc.add_table(rows=2, cols=2)
    table_geometry(frame_table, [4680, 4680])
    frame_spec = [
        ("02_revit_before.png", "01｜DWG–Revit 输入与轴网", "保留 Revit 属性栏和轴网区域；配合 07–34 秒。"),
        ("03_revit_after.png", "02｜房间结果与 IFC 参数", "先局部推近参数栏，后拉远展示彩色空间；配合 34–72 秒。"),
        ("01_agent_workflow.png", "03｜Agent 阶段结果", "仅保留房间创建、赋参、另存等结果行；正文滚动用 2–3 倍速。"),
        ("04_quality_report.png", "04｜SZ-IFC 报告输出", "保留“质检报告已生成”；项目路径与用户名打码。"),
    ]
    for cell, spec in zip([c for row in frame_table.rows for c in row.cells], frame_spec):
        filename, caption, note = spec
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(str(FRAME_DIR / filename), width=Inches(2.95))
        p = cell.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(2)
        font(p.add_run(caption), size=8.8, bold=True, color="1F4D78")
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        font(p.add_run(note), size=8.2, color="536779")

    h1(doc, "剪辑与配音执行要求")
    add_bullets(doc, [
        "每一句旁白必须落在对应技术画面上；不要在无关的软件等待界面讲算法。",
        "术语第一次出现时，用 4–8 字中文解释：例如“DXF 图元解析”“标准 IFC 标识库”“受限候选复核”。",
        "开篇定义一次“Transformer 架构 BIM Agent”；后续统一使用“Agent”“语义模块”“决策模块”等角色化表达。",
        "字幕分两层：底部完整旁白；左上角仅保留一个技术标签。画面停留 1 秒以上才放长字幕。",
        "公开发布前打码本地绝对路径、用户名和未经授权的项目编号。",
    ])

    doc.save(OUTPUT)
    with ZipFile(OUTPUT) as archive:
        assert archive.testzip() is None
        assert len([n for n in archive.namelist() if n.startswith("word/media/")]) == 4
    print(OUTPUT)


if __name__ == "__main__":
    build()
